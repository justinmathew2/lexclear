import io
import docx

def parse_docx(file_bytes: bytes) -> str:
    """Extract text from a Microsoft Word .docx file."""
    docx_file = io.BytesIO(file_bytes)
    doc = docx.Document(docx_file)
    extracted_text = []
    
    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            extracted_text.append(text)
            
    # Also extract tables if present
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                extracted_text.append(f"[Table Row] {row_text}")
                
    return "\n\n".join(extracted_text)
